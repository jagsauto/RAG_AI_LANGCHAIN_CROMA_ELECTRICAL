"""Create sample PDF, Excel, and Word files for electrical product data."""
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DOCUMENTS_DIR = PROJECT_ROOT / "data" / "documents"
PDFs_DIR = DOCUMENTS_DIR / "PDFs"
XLS_DIR = DOCUMENTS_DIR / "XLS"
TEXT_FILES_DIR = DOCUMENTS_DIR / "Text_Files"

for d in (PDFs_DIR, XLS_DIR, TEXT_FILES_DIR):
    d.mkdir(parents=True, exist_ok=True)

def create_pdf():
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        from reportlab.lib.units import mm
    except ImportError:
        print("Install reportlab to generate sample PDF: pip install reportlab")
        return
    path = PDFs_DIR / "product_sheet.pdf"
    c = canvas.Canvas(str(path), pagesize=A4)
    c.setFont("Helvetica-Bold", 16)
    c.drawString(40, 800, "Electrical Product Sheet – LED Bulbs & Cables")
    c.setFont("Helvetica", 11)
    c.drawString(40, 760, "LED Bulb E27 9W: 806 lm, 2700K, 15000h. Price: £4.99.")
    c.drawString(40, 740, "LED Bulb B22 12W: 1055 lm, 6500K, 20000h. Price: £5.49.")
    c.drawString(40, 710, "Twin & Earth 1.5mm² 100m: lighting circuits, 14A. £45.")
    c.drawString(40, 690, "Twin & Earth 2.5mm² 100m: sockets, 24A. £62.")
    c.drawString(40, 660, "Switches: 1-Gang 1-Way £2.49, Dimmer 400W £12.99.")
    c.save()
    print(f"Created {path}")

def create_excel():
    XLS_DIR.mkdir(parents=True, exist_ok=True)
    # XLSX
    try:
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Inventory"
        rows = [
            ["SKU", "Name", "Category", "Price", "Stock"],
            ["B-E27-9W", "LED Bulb E27 9W", "Bulbs", 4.99, 120],
            ["B-B22-12W", "LED Bulb B22 12W", "Bulbs", 5.49, 80],
            ["C-1.5-TE", "Twin & Earth 1.5mm² 100m", "Cables", 45.00, 25],
            ["C-2.5-TE", "Twin & Earth 2.5mm² 100m", "Cables", 62.00, 20],
            ["S-1G-1W", "1-Gang 1-Way Switch", "Switches", 2.49, 200],
            ["S-DIM-400", "Dimmer Switch 400W", "Switches", 12.99, 45],
        ]
        for r, row in enumerate(rows, 1):
            for c, val in enumerate(row, 1):
                ws.cell(row=r, column=c, value=val)
        path_xlsx = XLS_DIR / "inventory.xlsx"
        wb.save(str(path_xlsx))
        print(f"Created {path_xlsx}")
    except ImportError:
        print("Install openpyxl to generate sample XLSX: pip install openpyxl")
    # XLS (optional)
    try:
        import xlwt
        wb = xlwt.Workbook()
        ws = wb.add_sheet("Price List")
        for c, h in enumerate(["SKU", "Name", "Price"]):
            ws.write(0, c, h)
        for r, (sku, name, price) in enumerate([
            ("B-E27-9W", "LED Bulb E27 9W", 4.99),
            ("C-2.5-TE", "Twin & Earth 2.5mm²", 62.00),
        ], 1):
            ws.write(r, 0, sku)
            ws.write(r, 1, name)
            ws.write(r, 2, price)
        path_xls = XLS_DIR / "price_list.xls"
        wb.save(str(path_xls))
        print(f"Created {path_xls}")
    except ImportError:
        pass  # xlwt optional


def create_docx():
    try:
        from docx import Document as DocxDocument
        from docx.shared import Pt
    except ImportError:
        print("Install python-docx to generate sample DOCX: pip install python-docx")
        return
    path = TEXT_FILES_DIR / "product_guide.docx"
    doc = DocxDocument()
    doc.add_heading("Electrical Product Guide", 0)
    doc.add_paragraph(
        "LED Bulbs: E27 9W (806 lm, 2700K) £4.99. B22 12W (1055 lm, 6500K) £5.49. "
        "Cables: Twin & Earth 1.5mm² £45, 2.5mm² £62, 4mm² £95. "
        "Switches: 1-Gang 1-Way £2.49, Dimmer 400W £12.99."
    )
    doc.add_paragraph("All items suitable for domestic use. Cables in 100m drums.")
    doc.save(str(path))
    print(f"Created {path}")


if __name__ == "__main__":
    create_pdf()
    create_excel()
    create_docx()
